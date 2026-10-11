#!/usr/bin/env python3
"""Builds trackers/README.md (editor-friendly dashboard), docs/17-EXECUTION-TRACKER.md and
trackers/Delivery_Tracker.xlsx (+ CSVs).

Single source: the CORE backlog below + backlog tables parsed from the docs
(TPL-, UX-, INT-, CFG-, AI-, SEC-, T-ids). Re-run after editing any of them:
    python3 scripts/build_tracker.py
"""
import csv, re, pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule

ROOT = pathlib.Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"

MILESTONES = [  # phased releases (07 v0.3); old M-codes kept as work-package refs
 ("P0","Spikes + G1 + design system","2026-11-02","2026-11-20","G1","SP-01…SP-10, ADRs, tokens, CI + E2E harness smoke (old M0–M1)"),
 ("R1","Workspace & Policies","2026-11-23","2027-01-15","G-R1","System console, auth+MFA, roles, settings, DataGrid, SOP/policy library, campaigns, announcements, files, audit, dashboard (old M2–M4 part)"),
 ("R2","Requests & Approvals","2027-01-18","2027-03-12","G-R2","Runtime core, forms, designer, portal, inbox, delegation, SLA, 6 templates (old M5–M6 part)"),
 ("R3","Operate & Sign","2027-03-15","2027-05-07","G-R3","Admin console, DoA, master data, recall/edit, e-sign v1, 13 templates, PWA, BPMN import (old M4 rest, M6 rest, M7)"),
 ("R4","Enterprise Connect","2027-05-10","2027-07-02","G-R4","SSO/SCIM, scoped admins, webhooks/API/CLI, email approvals, guests, e-sign v2 (old M8)"),
 ("R5","Cases & Compliance","2027-07-05","2027-08-27","G-R5","Cases, investigations, legal hold, Part 11, reports, AI optional (old M9)"),
 ("GA","1.0 hardening + release","2027-08-30","2027-10-15","G4/G5","Load/HA/DR, pentest, a11y/VPAT, i18n, docs, full E2E regression (old M10–M11)"),
 ("POST","Post-GA expansion","2027-10-18","2028-03-31","Demand","Chat platforms, HRIS/ERP, QES/Aadhaar, BPMN/DMN, mining, AI agents (old M12+)"),
]
OLD_TO_REL = {"M0":"P0","M1":"P0","M2":"R1","M3":"R1","M4":"R1","M5":"R2","M6":"R2","M7":"R3","M8":"R4","M9":"R5","M10":"GA","M11":"GA","M12+":"POST"}
REL_OVERRIDE = {  # items whose release differs from the plain M→R mapping
 "BPMS-066":"R3","BPMS-102":"R3","BPMS-104":"R3","BPMS-107":"R3","BPMS-108":"R3","BPMS-109":"R3","BPMS-110":"R3",
 "BPMS-120":"R1","BPMS-121":"R2","BPMS-186":"R5",
}

# id, area, title, pri, milestone, owner_role, refs, acceptance, size, depends
CORE = [
 # M0/M1
 ("BPMS-001","DISC","User interviews (12+, startup/SME/MNC × HR/IT/Fin/Compliance/Admin) + synthesis","P0","M0","Product/UXR","OQ-05, A-*","Synthesis published; personas validated","M",""),
 ("BPMS-002","DISC","Model 5 reference processes in IR on paper","P0","M0","Runtime lead","A-02","All 5 expressible or IR gaps logged","S",""),
 ("BPMS-003","SPIKE","SP-01…SP-06 runtime spikes (native vs Temporal vs Operaton)","P0","M0","Runtime","U-01, ADR-0015","Reports in docs/spikes/","L",""),
 ("BPMS-004","SPIKE","SP-07 licence/SBOM review of candidate deps","P0","M0","Tech lead","ADR-0014","No blocked licence","S",""),
 ("BPMS-005","SPIKE","SP-08 expression language (dsntk FEEL vs CEL, incl. WASM)","P0","M0","Runtime","U-02, ADR-0010","Engine passes IR expression suite","M",""),
 ("BPMS-006","SPIKE","SP-09 Cedar + RLS + sqlx pool + partial eval","P0","M0","Platform","U-03, ADR-0011","0 leaks; list P95 in budget","M",""),
 ("BPMS-007","SPIKE","SP-10 audit chain 10M events: append + explorer","P0","M0","Platform","U-04","≥2k ev/s; P95<2s","M",""),
 ("BPMS-008","UX","Clickable prototype of Wave-1 flows + 5-user test","P0","M0","Design","05","Findings logged","M",""),
 ("BPMS-009","ARCH","G1: accept ADR-0010/0011/0015; freeze IR v1","P0","M1","Architect","G1","Signed off","S","BPMS-003..007"),
 ("BPMS-010","UX","Design tokens + core components + visual direction","P0","M1","Design/FE","05 §7, UX-001","axe clean; approved","L",""),
 ("BPMS-011","ENG","Monorepo, CI gates (fmt, clippy, deny, tests, a11y, SBOM)","P0","M1","DevOps","03 §11","All gates blocking","M",""),
 ("BPMS-012","SEC","Threat model review","P0","M1","Security","06 §2","Reviewed; tests mapped","S",""),
 # M2
 ("BPMS-020","PLAT","System plane: bootstrap, system roles, workspaces lifecycle, quotas","P0","M2","Platform","FR-S01–S04","J8 passes","L","BPMS-011"),
 ("BPMS-021","IAM","Local auth (Argon2id), email verify, lockout, sessions","P0","M2","Platform","12 §3, FX-I01","SEC tests green","M","BPMS-020"),
 ("BPMS-022","IAM","MFA TOTP + WebAuthn, recovery codes, step-up","P0","M2","Platform","12 §3/§11","MFA required by role works","M","BPMS-021"),
 ("BPMS-023","IAM","Memberships, invites, directory (departments/groups/positions/managers), CSV import","P0","M2","Platform","FR-W01/W02","Manager cycles rejected","M","BPMS-020"),
 ("BPMS-024","IAM","Permission catalogue, role templates, custom roles, grants, scopes, preview → Cedar","P0","M2","Platform","FR-W03, ADR-0011","No-escalation property test","L","BPMS-023"),
 ("BPMS-025","SEC","Forced RLS + repository WsCtx + CI cross-workspace sweep","P0","M2","Platform","03 §7, SEC-01/02","Sweep green in CI","M","BPMS-020"),
 ("BPMS-026","AUD","Audit chain append + coverage contract test","P0","M2","Platform","FR-T01, AUD-01/07","Every mutating route emits 1 event","M","BPMS-025"),
 ("BPMS-027","CFG","Settings layering, floors, profiles, explainer","P0","M2","Platform","FR-S05, FR-W04/W06, CFG-001..003","Never looser than floor","M","BPMS-020"),
 ("BPMS-028","UX","Shared DataGrid v1 (sort, filter builder, saved views, export, keyboard)","P0","M2","FE","FR-V01–V09","Used by Members & Workspaces","L","BPMS-010"),
 ("BPMS-029","ENG","Outbox + dispatcher + inbox dedup + idempotency keys","P0","M2","Platform","ADR-0005","T02 skeleton passes","M","BPMS-011"),
 ("BPMS-030","ENG","Key hierarchy: KeyProvider (sealed local, OpenBao, KMS), per-workspace KEK","P0","M2","Platform/Sec","04 §8, 06 §4","Rotation + shred tested","M","BPMS-020"),
 ("BPMS-031","UX","App shell, workspace switcher, system console shell, ⌘K","P0","M2","FE","UX-002/003","Keyboard nav","M","BPMS-010"),
 # M3
 ("BPMS-040","LIB","Artifacts, versions, lifecycle state machine, immutability trigger, SoD","P0","M3","Backend","FR-G04, T07","T07 passes","L","BPMS-024,026"),
 ("BPMS-041","LIB","Landscape, collections, owners, RACI, visibility","P0","M3","Fullstack","FR-G01/G02","Search authorized-only","M","BPMS-040"),
 ("BPMS-042","LIB","SOP/policy block editor + renderer (web/PDF)","P0","M3","FE","FR-G03, ADR-0013, U-10","Tagged PDF export","L","BPMS-040"),
 ("BPMS-043","LIB","Semantic diff across artifact kinds","P0","M3","Fullstack","FR-G05","Diff covers bundle artifacts","M","BPMS-040"),
 ("BPMS-044","LIB","Reviews, change requests, effective-date scheduling, renewal date, verification badge","P0","M3","Fullstack","FR-G05/G07, FX-G01/G02","Expired flagged","M","BPMS-040"),
 ("BPMS-045","LIB","Policy library + employee policy portal + applicability","P0","M3","Fullstack","FR-PL01–PL03","Applicability enforced","M","BPMS-042"),
 ("BPMS-046","CAMP","Campaign engine: ack/quiz-lite, audience snapshot, reminders, escalation, new joiners, evidence export","P0","M3","Backend","FR-AK01–04, ADR-0017","5k recipients/min","L","BPMS-029"),
 ("BPMS-047","COMMS","Announcements: compose, targeting, schedule/expiry, pin, required ack, approval, analytics","P0","M3","Fullstack","FR-AN01–06, FX-C01","J14 passes","L","BPMS-046"),
 ("BPMS-048","NOTIF","In-app notification centre + SMTP email + templates + digests","P0","M3","Fullstack","FR-N01/N02/N04","Bounce handling","M","BPMS-029"),
 ("BPMS-049","SRCH","Postgres FTS + trgm, permission-filtered search","P0","M3","Backend","04 §6","No restricted leakage","M","BPMS-040"),
 # M4
 ("BPMS-060","FILE","Files pipeline: resumable upload, hash, sniff, ClamAV, quarantine, signed URLs","P0","M4","Backend","FR-F01–F03, F05","EICAR quarantined; SEC-10","L","BPMS-030"),
 ("BPMS-061","FILE","Previews (image/PDF) in sandbox; Office preview spike","P0","M4","Backend","FR-F04, U-08","No active content","M","BPMS-060"),
 ("BPMS-062","FILE","Workspace Files view, quotas, orphan cleanup","P0","M4","FE","FR-F06/F07","Held files undeletable","M","BPMS-060,028"),
 ("BPMS-063","ADMIN","Workspace overview dashboard + people & access monitor","P0","M4","FE","FR-A01/A06","Every tile drills down","L","BPMS-028"),
 ("BPMS-064","AUD","Audit explorer, timelines, saved queries, export","P0","M4","Fullstack","FR-T03/T04","P95<2s @10M","L","BPMS-026"),
 ("BPMS-065","PLAT","System console: workspaces grid, quotas, health, system audit","P0","M4","FE","FR-S02–S09","—","M","BPMS-020"),
 ("BPMS-066","AUD","Retention policy per record class (policy only)","P0","M4","Backend","PRV-04","Jobs honour classes","M","BPMS-026"),
 # M5
 ("BPMS-080","RT","IR v1 model + compiler (§11 checks, diagnostics, semantic diff)","P0","M5","Runtime","11, ADR-0003","All checks covered","L","BPMS-009"),
 ("BPMS-081","RT","Expression engine integration (typed, budgeted, WASM build)","P0","M5","Runtime","ADR-0010","Budget exhaustion = incident","M","BPMS-009"),
 ("BPMS-082","RT","Interpreter: advance loop, tokens, splits/joins, end/terminate","P0","M5","Runtime","04 §4","Conformance green","L","BPMS-080"),
 ("BPMS-083","RT","Jobs + timers + scheduler with leases/fencing; calendars; SLA clocks","P0","M5","Runtime","04 §4.4, FR-E12","T06; lateness SLO","L","BPMS-082"),
 ("BPMS-084","RT","Release bundles + publish integration","P0","M5","Runtime","ADR-0012","T03 passes","M","BPMS-040,080"),
 ("BPMS-085","WORK","Work items, assignment strategies, manager chain, DoA matrix, quorum, per-group","P0","M5","Backend","FR-E04/E05, FX-D01/D02, FX-E01","T01 passes","L","BPMS-082,023"),
 ("BPMS-086","WORK","Delegation, OOO substitution, SoD re-check, send-back","P0","M5","Backend","12 §9/§10, FX-E05","T05 passes","M","BPMS-085"),
 ("BPMS-087","RT","Rule tables (UNIQUE/FIRST/COLLECT)","P0","M5","Runtime","FR-E10","Deterministic","M","BPMS-081"),
 ("BPMS-088","RT","Incidents, retry/skip/cancel, instance admin ops","P0","M5","Backend","FR-E13","One open incident per failure","M","BPMS-082"),
 ("BPMS-089","QA","Conformance suite + model-based tests + deterministic simulation harness","P0","M5","SDET/Runtime","15 §4–6","Nightly chaos green","L","BPMS-082"),
 # M6
 ("BPMS-100","DES","Step builder + React Flow canvas + properties + validate/test-run/review/publish","P0","M6","FE","FR-E01/E11, ADR-0018","T10 keyboard-only","XL","BPMS-080"),
 ("BPMS-101","FORM","Form builder + renderer (calc, prefill, repeating, lookups, PDF of submission)","P0","M6","FE","FR-E03, FX-F01–F04","Server = client validation","XL","BPMS-081"),
 ("BPMS-102","DATA","Master-data tables (CSV import, lookups, versioned)","P0","M6","Fullstack","FX-E02","Cascading dropdowns","M","BPMS-101"),
 ("BPMS-103","REQ","Request portal/catalogue, start form, request detail, comments/@mentions/watchers","P0","M6","Fullstack","FR-R01/R02, FX-R01/R02","J0/J3 pass","L","BPMS-082"),
 ("BPMS-104","REQ","Recall/withdraw, edit-after-submit, request-on-behalf, numbering","P0","M6","Backend","FX-E03/E04/E06","Audited","M","BPMS-103"),
 ("BPMS-105","WORK","My Work 3-pane inbox (tasks, acks, signatures), bulk with preview","P0","M6","FE","FR-R03, FR-AK02","Decision ≤3 interactions","L","BPMS-085,046"),
 ("BPMS-106","TPL","Template gallery + install wizard + Wave-1 templates (13)","P0","M6","Fullstack/SMEs","FR-G08, TPL-001..","All W1 test-run green","XL","BPMS-100,101"),
 ("BPMS-107","SIGN","E-sign v1: PDF, recipients/roles, routing, bulk, fields, consent step, ceremony, certificate","P0","M6","Fullstack","FR-DS01–07, FX-S01/02/05/06/07","J13 passes","XL","BPMS-046,060"),
 ("BPMS-108","ADMIN","All requests / all work / workload / process health / ops monitor","P0","M6","FE","FR-A02–A05","Bulk audited","L","BPMS-063,085"),
 ("BPMS-109","UX","PWA (360px) for My Work, requests, acks, signing","P0","M6","FE","FR-P07","Lighthouse PWA pass","M","BPMS-105"),
 ("BPMS-110","LIB","BPMN 2.0 XML import (subset) + compatibility report","P0","M6","Runtime","FR-G09, SP-06","Unsupported listed","M","BPMS-080"),
 # M7
 ("BPMS-120","PILOT","Pilot workspaces (Starter + Enterprise profile), 3–5 real processes, 4 weeks","P0","M7","Product","G3","G3 checklist","L","M6"),
 ("BPMS-121","UX","Usability round + SUS measurement + fixes","P0","M7","Design","02 §9","SUS targets","M","BPMS-120"),
 # M8
 ("BPMS-140","IAM","OIDC SSO, SAML-broker, SSO enforcement, JIT","P1","M8","Platform","FR-P02, INT-001/002","Deprovision test","L",""),
 ("BPMS-141","IAM","SCIM 2.0 server","P1","M8","Platform","INT-003","<5 min deprovision","M","BPMS-140"),
 ("BPMS-142","IAM","Scoped/delegated admins; IP allowlists","P1","M8","Platform","CFG-018, FR-W06","No escalation","M","BPMS-024"),
 ("BPMS-143","INT","Webhooks (Standard Webhooks), HTTP connector via egress proxy, inbound events","P1","M8","Backend","INT-*, WH/SSRF","SEC SSRF tests","L","BPMS-029"),
 ("BPMS-144","INT","Public API hardening, OpenAPI, CLI, TS/Python SDKs","P1","M8","Backend","08","Spec published","M",""),
 ("BPMS-145","NOTIF","Email-actionable approvals, reply-by-email, web push","P1","M8","Fullstack","FR-N03/N05, FX-R03","Replay rejected","M","BPMS-048"),
 ("BPMS-146","IAM","Guests + external signers (magic link + OTP); self-serve signing links","P1","M8","Fullstack","FR-P08, FX-S10","Single-use","M","BPMS-107"),
 ("BPMS-147","SIGN","E-sign v2: DOCX + merge, anchor tags, conditional fields, delegate, correction, embedded","P1","M8","Fullstack","FX-S03/04/08/09/12","Audited corrections","L","BPMS-107"),
 ("BPMS-148","FORM","Public forms (privacy CAPTCHA), save & resume, partials, form signature field","P1","M8","Fullstack","FX-F05–F07","Rate-limited","M","BPMS-101"),
 ("BPMS-149","EXEC","Subprocesses, reusable sub-flows, duplicate detection, bulk import, env promotion","P1","M8","Runtime","FR-E08, FX-E07–E09, FR-P09","—","L","BPMS-082"),
 ("BPMS-150","LIB","Auto-generated process documentation; policy exceptions; policy templates","P1","M8","Fullstack","FX-G04, FR-PL04, FX-G03","—","M","BPMS-042"),
 ("BPMS-151","COMMS","Emergency alerts, content validation period, moderation, quiz","P1","M8","Fullstack","FX-C02–C04","—","M","BPMS-047"),
 # M9
 ("BPMS-160","CASE","Case types, queues, macros, SLAs, link/merge, restricted cases","P1","M9","Fullstack","FR-C01, FX-K01–K06","J4 passes","XL","BPMS-082"),
 ("BPMS-161","AUD","Investigations: pins, legal hold, signed export, conflict-of-interest","P1","M9","Fullstack","FR-T05/T06, FX-K08, FX-L01","J10 passes","L","BPMS-064,160"),
 ("BPMS-162","AUD","Checkpoints, RFC 3161, WORM export, procera-verify CLI","P1","M9","Platform","AUD-02–05","Verifier detects tamper","M","BPMS-026"),
 ("BPMS-163","GOV","Attestation campaigns, control mapping, access reviews","P1","M9","Fullstack","FR-G06, FR-W11","Evidence export","M","BPMS-046"),
 ("BPMS-164","GOV","Retention disposition + certificates, DSAR template, legal hold engine","P1","M9","Backend","PRV-01/03/04","Holds block deletion","M","BPMS-066"),
 ("BPMS-165","SIGN","Part 11 mode, PAdES-LTV, TSA; in-person/kiosk signing","P1","M9","Fullstack","ESIG-01–05, FR-DS05, FX-S11","Verification passes","L","BPMS-107"),
 ("BPMS-166","ADMIN","Alert centre, custom dashboards, report builder, scheduled reports, KPI goals, usage analytics, CSAT","P1","M9","Fullstack","FR-A07–A09, FX-A01–A04, FX-R04","RLS at run time","XL","BPMS-063"),
 ("BPMS-167","IAM","Separation mode; break-glass","P1","M9","Platform","12 §6.1, FR-S07","Visible to workspace","M","BPMS-024"),
 ("BPMS-168","LIB","Role-based onboarding/training paths","P1","M9","Fullstack","FX-G07","Completion reports","M","BPMS-046"),
 # M10/M11
 ("BPMS-180","OPS","Load test at reference load; HA failover; RPO/RTO restore drill","P1","M10","SRE","06, T09","Targets met","L",""),
 ("BPMS-181","SEC","External pentest + ASVS L2 self-assessment + fixes","P1","M10","Security","06","0 crit/high","M",""),
 ("BPMS-182","UX","WCAG 2.2 AA audit + VPAT 2.5 INT","P1","M10","Design/QA","A11Y-01/03","Audit pass","M",""),
 ("BPMS-183","I18N","3 UI languages + RTL + pseudo-locale","P1","M10","FE","FR-P04, OQ-09","Checks pass","M",""),
 ("BPMS-184","OPS","Helm chart, air-gapped bundle, upgrade N-1→N tests, dedicated-DB tier","P1","M10","DevOps","FR-P06, FR-S10","Upgrade with live instances","L",""),
 ("BPMS-185","DOC","Docs site, admin guide, API reference, trust pack","P1","M10","Writer","FX-P06","Published","M",""),
 ("BPMS-186","AI","AI assistance P1 set (drafting, summaries, translate, Q&A) behind kill switch","P1","M10","Fullstack","16, ADR-0006","Eval thresholds","L",""),
 ("BPMS-190","REL","GA release: signed artefacts, SBOM, provenance, runbooks","P1","M11","Tech lead","G5","Released","S","M10"),
 # M12+
 ("BPMS-200","INT","Chat/host-platform connectors (identity, embed, cards, notifications, directory)","P2","M12+","Backend","13 §4","Demand gate","XL",""),
 ("BPMS-201","INT","Chat cards (Slack/Teams), M365/Google, HRIS, ERP connectors","P2","M12+","Backend","INT-*","—","XL",""),
 ("BPMS-202","SIGN","QES via QTSP; Aadhaar eSign via ESP; SMS/IDV providers","P2","M12+","Backend","FX-S13, FX-S07","—","L",""),
 ("BPMS-203","EXEC","BPMN view (diagram-js), DMN import/export, migration tool, simulation","P2","M12+","Runtime/FE","ADR-0018, FR-E07/E10, FR-I04","—","XL",""),
 ("BPMS-204","INS","Event-log export OCEL 2.0 / XES","P2","M12+","Data","FR-I03","Schema-valid","M",""),
 ("BPMS-205","AI","Governed AI agent steps","P2","M12+","Fullstack","16","Human gate","L",""),
 ("BPMS-013","E2E","E2E harness: Playwright + compose.e2e (PG, SeaweedFS, Valkey, Mailpit, ClamAV), seed A/B, personas, __test clock/drain/fault (feature-gated)","P0","M1","SDET","23 §2–4","Smoke runs in CI; endpoints absent in release build","M","BPMS-011"),
 ("BPMS-032","E2E","E2E-R1 pack (R1-01…19 + ISO)","P0","M4","SDET + feature owners","23 §5","Green on Chromium + mobile","L","BPMS-013"),
 ("BPMS-111","E2E","E2E-R2 pack (R2-01…16 + ISO) + Firefox","P0","M6","SDET + feature owners","23 §5","Green; race/restart/timer stable","L","BPMS-032"),
 ("BPMS-112","E2E","E2E-R3 pack (R3-01…12 + MOB + ISO) + WebKit","P0","M7","SDET + feature owners","23 §5","Green; bulk-sign 1000","L","BPMS-111"),
 ("BPMS-152","E2E","E2E-R4 pack incl. Keycloak, SCIM, webhook receiver","P1","M8","SDET + feature owners","23 §5","Green","L","BPMS-112"),
 ("BPMS-169","E2E","E2E-R5 pack incl. investigations, Part 11, break-glass","P1","M9","SDET + feature owners","23 §5","Green","L","BPMS-152"),
 ("BPMS-187","E2E","E2E-GA regression: 3 browsers × 2 viewports, cache-off, upgrade-path R1→1.0, air-gapped","P1","M10","SDET","23 §5","Green; flake < 0.5 %","M","BPMS-169"),
 ("BPMS-014","ENG","AGENTS.md/CLAUDE.md contract, PR template (spec id + E2E id), CODEOWNERS for protected areas","P0","M1","Tech lead","07 §3","In repo","S",""),
 ("BPMS-034","BRAND","Brand vector masters (SVG symbol/wordmark/lockups, small-size symbol), favicon/PWA icon set from vector, trademark & domain search","P0","M2","Design + Product","assets/brand, OQ-01, U-16","SVGs in assets/brand/vector; clearance memo","S",""),
 ("BPMS-033","OPS","Valkey cache port (rate limit, session cache, policy/settings cache, SSE fan-out) + cache-off fallback","P0","M2","Platform","ADR-0019","E2E passes with cache off","M","BPMS-020"),
]
CORE = [r[:4]+(REL_OVERRIDE.get(r[0], OLD_TO_REL.get(r[4], r[4])),)+r[5:] for r in CORE]
CORE.sort(key=lambda r: ([m[0] for m in MILESTONES].index(r[4]), r[0]))

# Item status lives here (docs/24 §2.2). Default is "Not started".
# Values: Not started | In progress | Blocked | In review | Done | Deferred.  Format: id: (status, note)
STATUS = {
 "BPMS-002": ("Done", "5 reference processes as IR v1 fixtures (all 16 node kinds) + JSON Schema + structural checker IR-E001..E016 with negative tests; 12 IR spec clarifications (docs/spikes/BPMS-002)"),
 "BPMS-005": ("Done", "SP-08: recommend CEL + decimal/stdlib/type-checker/budgets; FEEL rejected at runtime (quadratic loops, no WASM) — docs/spikes/SP-08; ADR-0010 recommended"),
 "BPMS-006": ("Done", "SP-09: Cedar 6/6 correct, 5.6µs P50; RLS via pool 0 leaks/5000; LLD fix NULLIF in RLS policy; partial-eval→SQL viable with fallback — docs/spikes/SP-09"),
 "BPMS-007": ("Done", "SP-10: head-lock 9.7k ev/s (P99 11ms), batched sealer 65k ev/s, explorer P95 <0.6ms at 10M, verifier 536k ev/s + tamper detected — docs/spikes/SP-10"),
 "BPMS-003": ("Done", "SP-01..06 native runtime: 10k instances under 1329 crashes 0 lost; race exactly-once; join once; 410 retries 0 dup effects; 30-day pinning; BPMN 54-node lossless round trip — docs/spikes/SP-01-06; ADR-0015 recommended"),
 "BPMS-004": ("Done", "SP-07: Rust + JS dependency licences reviewed; approved set; JS licence gate (scripts/js-licence-gate.mjs) + CDLA-Permissive-2.0 for CA roots; fonts OFL — docs/spikes/SP-07"),
 "BPMS-010": ("Done", "Tokens (41 contrast checks + forbidden-pair guard), Button/TextField/Checkbox/Notice/Select/ConfirmDialog/Tabs/DataGrid shell on React Aria + TanStack, app shell, gallery; axe light/dark, keyboard, focus trap, aria-sort, reflow, targets (E2E-P0-05..13)"),
 "BPMS-014": ("Done", "CLAUDE.md, AGENTS.md, PR template, CODEOWNERS, docs/24 rules on main"),
 "BPMS-020": ("Done", "procera-platform + /sysapi/v1: one-time bootstrap token (printed at startup until bootstrapped), system.owner/admin/operator/auditor fixed authz table, MFA required for system plane, workspace create/list/get/rename/suspend/restore/schedule-deletion with If-Match + ETag, profile quotas, platform setting floors seeded, all audited; E2E-R1-SYS-01..05"),
 "BPMS-021": ("Done", "procera-identity + /api/v1/auth: users, Argon2id (policy, rehash, enumeration-safe), JWT access (ADR-0020) + rotating refresh with reuse detection, DB login throttling + per-IP rate limit, logout, sign-out-everywhere, deactivation via token_version; audited; E2E-R1-AUTH-01..05,07"),
 "BPMS-022": ("In progress", "TOTP (RFC 6238, replay-protected, sealed secrets) + 10 recovery codes + MFA login (E2E-R1-AUTH-06). Step-up: POST /session/step-up (TOTP code, or password without MFA) re-issues the token with a fresh auth_time; enforced against auth.stepup.window (401 step_up_required) for step-up settings, role create/edit and grant/revoke (E2E-R1-CFG-03). MFA required by role (auth.mfa.required_roles) enforced at workspace entry, 403 mfa_required (E2E-R1-MEM-10). Remaining: WebAuthn/passkeys"),
 "BPMS-030": ("In progress", "procera-crypto: LocalKey (env root key, HMAC-derived purpose keys), AES-256-GCM seal/open bound to context, token helpers, procera keygen. Remaining: KMS/OpenBao adapters, per-workspace KEKs, rotation"),
 "BPMS-023": ("Done", "procera-members + /api/v1: memberships, invitations (hashed + sealed tokens, 7-day expiry, email-bound, single use; list, revoke, resend), register-by-invite, owner onboarding activates the workspace, workspace switch (ws-scoped JWT), deactivation guards, state gate (423). Directory (migration 0008): department tree (uuid[] paths, sibling-unique names, no moves under descendants), static groups, positions (behind directory.positions.enabled), manager links with cycle rejection, deterministic manager resolution (explicit → nearest department head), per-field source attribution, CSV import with plan/apply. All audited; E2E-R1-MEM-01..03,05..07, E2E-R1-DIR-01..04"),
 "BPMS-024": ("Done", "procera-access: permission catalogue + 12 role templates; custom roles (no escalation on create/edit/grant/revoke, built-ins immutable, clone lineage); grants to members or static groups (migration 0009), scoped (workspace/department/collection/process/case) with expiry, compiled to Cedar; members inherit their department subtree so department-scoped admins reach only it; invariants (workspace isolation, active members only); last-owner protection; role change preview; access_revisions cache; all changes audited. Property tests: no escalation (64 cases) and preview matches effect (48 cases), both against Postgres. E2E-R1-MEM-04,08,09"),
 "BPMS-040": ("Done", "procera-library (migration 0012): artifacts + versions for SOPs and policies; lifecycle draft → in_review → approved → scheduled/published → deprecated → retired with guarded transitions; one open draft per artifact; reviews in numbered rounds, approvals from governance.review.* settings, author never reviews own version, distinct approvers, comment required to reject; content_hash = SHA-256(JCS) frozen at approval and a trigger forbids changing approved content or moving back (T07); publish supersedes, future effective_at schedules (worker + drain), reason for change on later versions, retire; step-up for controlled approve/publish; outbox artifact.published.v1; every transition audited. E2E-R1-LIB-01..05. Process artifacts arrive with the designer/compiler"),
 "BPMS-031": ("In progress", "Web app on the Notion/Attio direction (owner choice): warm paper palette in tokens (contrast-checked), full-bleed pages with a warm icon sidebar, workspace + account menus (persisted theme), action-first Home (acknowledgements, reviews, announcements with inline actions, library table, activity, people), Library (filters, Attio-style table, New document dialog), document reading view with properties, lifecycle actions (edit, submit, approve, request changes, publish with step-up, new version), acknowledgement callout, Announcements feed with read/ack, Members with invite dialog; split sign-in with product preview; single-flight token refresh. E2E-R1-UI-01..10 with axe. Remaining: ⌘K command palette (search API is ready), system console shell, mobile bottom bar"),
 "BPMS-043": ("In progress", "API done (UI pending: compare view + UI E2E). procera-library diff module: block-level LCS alignment of SOP/policy content with word-level diffs inside changed blocks (removals and additions paired by position within an unmatched run, same block type → changed); GET /versions/{a}/diff/{b} (same artifact only); unit + property test (applying a diff always rebuilds the target) + E2E-R1-LIB-06. Process/form/rule bundle artifacts diff once their editors land (BPMS-050+)."),
 "BPMS-044": ("In progress", "API done (UI pending: governance panel, change requests + UI E2E). Migration 0016: review cadence (per-artifact override or governance.review_cadence.default/.controlled), next_review_at set on every publish, recertification, worker flag of passed review dates (once; artifact.review_overdue.v1 → owner in-app notification), verification badge with governance.verification.validity expiry, change requests (open → accepted/declined/implemented, linked to the implementing version); all audited and in the audit contract; E2E-R1-LIB-07..10. Effective-date scheduling shipped with BPMS-040."),
 "BPMS-045": ("In progress", "API done (UI pending: Policies portal page, applicability control + UI E2E). Migration 0017: policy category, applicability (audience; NULL = everyone) and restricted flag; procera_audience_contains SQL twin of the audience resolver; reader visibility (A-14): view-only readers see published history only, restricted policies hidden outside their applicability in lists, reads, diff, change requests and search; PUT /artifacts/{id}/applicability (audited); GET /me/policies portal with what-changed summaries vs the previous published version, text + category filters; E2E-R1-POL-01..03. Remaining: re-acknowledgement on major publish (FR-PL03), superseded watermark in exports."),
 "BPMS-046": ("Done", "procera-campaigns (migration 0013): acknowledgement campaigns for a published SOP/policy version; audience via the shared resolver (snapshot or dynamic), assignments with in-app + email notices, acknowledgement bound to version id + content hash + statement + auth method, reminder ladder and one-time manager escalation by the scheduler, dynamic new joiners, close/cancel, progress, audited evidence CSV; settings campaign.*; E2E-R1-CMP-01..07. Later: quiz-lite, exemptions, signed PDF evidence, auto re-ack on major publish"),
 "BPMS-047": ("Done", "procera-announcements (migration 0014): drafts with block content, audience snapshot, approval per announcements.approval_required with SoD, scheduling + derived expiry, pinned/critical, channels from settings, in-app + email delivery via the outbox, read and acknowledgement receipts, analytics, archive; settings announcements.*; E2E-R1-ANN-01..06. Later: poster scoping, reminders/escalation for unacked, comments/reactions, reach preview"),
 "BPMS-049": ("Done", "procera-search (migration 0015): Postgres FTS (english + simple) + pg_trgm fuzzy/prefix over SOPs, policies (artifact tsv kept current by triggers), members, departments, groups; one ranked list with escaped <mark> snippets; permission-filtered per type with Cedar before querying; RLS isolation; E2E-R1-SRCH-01..03"),
 "BPMS-064": ("In progress", "API done: procera_audit::explorer — filtered newest-first keyset paging, view-own scope, member names, artifact-aware object timelines, chain verify, private audited saved queries, capped CSV/JSONL exports with step-up that are themselves audited; E2E-R1-AUDX-01..06. Remaining: explorer UI (docs/05 §3.19), async export job, POST /audit/query"),
 "BPMS-048": ("In progress", "procera-notify (migration 0011): templates in code (minijinja, strict, HTML auto-escaped), in-app inbox (list/unread/read/read-all), email queue with sealed bodies and dedupe, SMTP delivery via lettre (STARTTLS/TLS) with backoff (max 1 h, 12 tries), workspace sender name; outbox consumers: invitation email with accept link, member-joined in-app to the inviter. Flags --public-url, --smtp-url, --mail-from. E2E-R1-NOT-01..02. Remaining: digests + quiet hours (FR-N04), localisation + terminology (CFG-008), more event kinds as features land"),
 "BPMS-025": ("Done", "procera-db: migrations (0001-0003), roles bootstrap (app/migrator/discovery), WsTx with tx-local set_config, procera_enable_workspace_rls() with NULLIF, procera_rls_violations guard (tested + enforced at startup), discovery.* functions with column grants; 6 DB tests"),
 "BPMS-026": ("Done", "procera-audit: JCS/SHA-256 per-workspace chain in the business tx, batched sealer, verifier, monthly partitions with forced RLS, append-only grants + trigger; tamper test. Route audit contract (audit_contract.rs): a unit test fails on any mutating route without an entry; E2E-R1-AUD-01 checks every route's appended events against it. Gaps closed: TOTP enrolment start, workspace selection, role/grant changes. data_access events follow with restricted data (BPMS-040+)"),
 "BPMS-027": ("In progress", "procera-settings: catalogue in code (P0 subset, 38 keys) with comparators; platform floors + workspace values + append-only revisions (migration 0007, RLS); 4 configuration profiles applied at workspace creation; resolution with clamping when a floor is raised; explainer; /api/v1/settings (list, explainer, PUT per key with If-Match, history) and /sysapi/v1/settings floors; reasons for ⚡ keys, step-up = MFA session (interim); audited (setting.changed, platform.floor_changed); property test (512 cases, never looser than floor); E2E-R1-CFG-01..05. Remaining: department/collection overrides (needs directory), profile switch plan/apply (CFG-017), settings cache invalidation (CFG-004)"),
 "BPMS-029": ("Done", "Outbox/inbox/idempotency tables + helpers (SKIP LOCKED claims, replay/conflict). Dispatcher (procera_db::dispatch, migration 0010): due-workspace discovery, per-workspace advisory lock, per-aggregate ordered delivery to a Bus, at-least-once with exponential backoff (max 5 min) and hold-back behind a failed event, NOTIFY trigger wake-ups + 1 s polling; runs in serve --all; /__test/drain really drains. Tests: order/backoff/hold-back/workspaces + NOTIFY wake-up; E2E-R1-MEM-07. Consumers (mailer) arrive with BPMS-048"),
 "BPMS-033": ("Done", "procera-cache: memory + Valkey/Redis (redis 1.7, BSD-3) with ws:/sys: key namespacing, TTL-only writes, Lua fixed-window rate limit with in-process fallback, degraded health; wired into procera (/readyz), CI Valkey service"),
 "BPMS-034": ("In progress", "Raster assets + brand guide merged; SVG masters and trademark search pending"),
 "BPMS-013": ("Done", "e2e feature + /__test controls, Playwright harness, P0 smoke pack (E2E-P0-01..04, REL), compose.e2e stack, CI e2e-smoke job; on main, CI green. Seed A/B + personas land with the first workspace APIs (BPMS-032); full compose boot pending a Docker run"),
 "BPMS-011": ("Done", "Rust workspace, kernel, procera binary, CI (fmt/clippy/test/deny/SBOM/E2E smoke) + web check/build, JS licence gate and axe a11y job (added with BPMS-010)"),
}
def status_of(i): return STATUS.get(i, ("Not started", ""))

RISKS = [
 ("R1","Execution correctness","Critical","High","IR spec, conformance, property/model tests, chaos","Runtime lead"),
 ("R2","Scope creep (BPMN/features)","High","High","Non-goals, catalogue change process, scope-cut list","Product"),
 ("R3","Cross-workspace data leak","Critical","Medium","Layered isolation, CI sweep, pentest","Security"),
 ("R4","Adoption / UX quality","High","Medium","Templates, Starter profile, usability per milestone, SUS gate","Design"),
 ("R5","Wrong runtime choice","High","Medium","Spikes before G1, adapter seam","Architect"),
 ("R6","Licence contamination","High","Medium","CI licence gate, SP-07 quarterly","Tech lead"),
 ("R7","Instance migration demand","High","High","Pinning; cancel-and-restart; P2 tool","Runtime"),
 ("R8","Duplicate external side effects","High","High","Outbox, idempotency keys","Backend"),
 ("R9","Analytics inaccuracy","Medium","Medium","Event semantics, reconciliation tests","Data"),
 ("R10","Staffing below plan","High","Medium","Option B timeline; scope cuts","Product"),
 ("R11","E-sign legal validity per jurisdiction","Medium","Medium","Legal posture doc, certificate, QES/Aadhaar providers P2","Legal"),
 ("R12","Admin visibility vs confidential cases","Medium","Medium","Restricted cases + Separation mode (OQ-12)","Product"),
 ("R13","Scanner/converter/object store availability","Medium","Low","Fail-closed + degraded modes","SRE"),
 ("R14","Email deliverability for self-hosters","Medium","Medium","Guides, bounce handling, in-app fallback","Backend"),
 ("R15","Accessibility regressions in designer/canvas","Medium","Medium","Keyboard alternative, axe in CI, audits","Design"),
 ("R16","AI-generated code quality / subtle bugs","High","Medium","Spec-first tasks, small PRs, CI gates, 2-human review on protected areas, E2E per feature","Tech lead"),
 ("R17","E2E flakiness slows releases","Medium","Medium","Test clock + drain, seeded data, 1 retry max, quarantine lane, flake budget","SDET"),
]

DECISIONS = []
for p in sorted((DOCS/"adr").glob("ADR-*.md")):
    t = p.read_text().splitlines()
    title = t[0].split("—",1)[1].strip()
    status = re.search(r"Status: \*\*(.+?)\*\*", p.read_text()).group(1)
    DECISIONS.append((p.stem[:8], title, status, p.name))

def parse_ids(pattern, files=None, maxcols=None):
    rows=[]
    for f in (files or sorted(DOCS.glob("*.md"))):
        for line in f.read_text().splitlines():
            if re.match(r"^\|\s*\**"+pattern, line):
                cells=[c.strip().strip("*") for c in line.strip().strip("|").split("|")]
                rows.append([f.name]+cells[:maxcols] if maxcols else [f.name]+cells)
    return rows

DOMAIN = {
 "Templates (TPL)": parse_ids(r"TPL-\d"),
 "UX (UX)": parse_ids(r"UX-\d"),
 "Integrations (INT)": parse_ids(r"INT-\d", [DOCS/"13-INTEGRATIONS.md"]),
 "Config (CFG)": parse_ids(r"CFG-\d", [DOCS/"20-WORKSPACE-CONFIGURATION.md"]),
 "AI (AI)": parse_ids(r"AI-\d", [DOCS/"16-AI-ASSISTANCE.md"]),
 "E2E scenarios (E2E)": parse_ids(r"E2E-(?:R\d|P0)-", [DOCS/"23-E2E-TEST-PLAN.md"]),
}
SECTESTS = parse_ids(r"SEC-\d", [DOCS/"06-SECURITY-AND-OPERATIONS.md"])
TESTS = parse_ids(r"T\d\d\b", [DOCS/"15-TEST-STRATEGY-AND-CONFORMANCE.md"])

# open questions / unknowns / assumptions from doc 22
def table_after(fname, prefix):
    return parse_ids(prefix, [DOCS/fname])
OQS = table_after("22-UNKNOWNS-ASSUMPTIONS-AND-OPEN-QUESTIONS.md", r"OQ-\d")
UNK = table_after("22-UNKNOWNS-ASSUMPTIONS-AND-OPEN-QUESTIONS.md", r"U-\d")
ASM = table_after("22-UNKNOWNS-ASSUMPTIONS-AND-OPEN-QUESTIONS.md", r"A-\d")

RACI = [
 ("Activity","Process Owner","Designer","Reviewer/Compliance","Workspace Admin","System Admin"),
 ("Draft process/SOP/policy","A","R","C","I","—"),
 ("Approve controlled change","C","I","A/R","I","—"),
 ("Publish standard process","A/R","C","C","I","—"),
 ("Configure roles & settings","I","—","C","A/R","—"),
 ("Install template","A","R","C","R","—"),
 ("Configure integration secret","I","C","C","A/R","C (allowed list)"),
 ("Run attestation campaign","C","—","A/R","R","—"),
 ("Investigate audit event","I","—","A/R","R","—"),
 ("Create/suspend workspace","—","—","—","I","A/R"),
 ("Approve break-glass","—","—","C","A","R (requester)"),
]

def write_md():
    out=["# 17 — Execution tracker","",
     "Status: Living · Owner: Product + Tech lead · **Generated** by `scripts/build_tracker.py` (edit the script's CORE list or the domain backlog tables in the docs, then re-run). Spreadsheet: [`trackers/Delivery_Tracker.xlsx`](../trackers/Delivery_Tracker.xlsx).","",
     "Release = phased release in 07 v0.3 (domain backlogs below still cite old M-codes; map them with: M0–M1→P0, M2–M4→R1, M5–M6→R2, M7→R3, M8→R4, M9→R5, M10–M11→GA, M12+→POST, except items moved per 07 §2). Status values: `Not started · In progress · Blocked · In review · Done · Deferred`. Size: S ≤ 3 d · M ≤ 2 wk · L ≤ 4 wk · XL > 4 wk (split before start).","",
     "## Releases (phased; see 07)","","| Release | Name | Start | End | Gate | Scope |","|---|---|---|---|---|---|"]
    out+= [f"| {m[0]} | {m[1]} | {m[2]} | {m[3]} | {m[4]} | {m[5]} |" for m in MILESTONES]
    out+=["","## Core backlog (epics)","","| ID | Area | Work item | Pri | Release | Owner | Refs | Acceptance | Size | Depends | Status |","|---|---|---|---|---|---|---|---|---|---|---|"]
    out+= [f"| {r[0]} | {r[1]} | {r[2]} | {r[3]} | {r[4]} | {r[5]} | {r[6]} | {r[7]} | {r[8]} | {r[9]} | {status_of(r[0])[0]} |" for r in CORE]
    from collections import Counter
    cnt=Counter(status_of(r[0])[0] for r in CORE)
    out+=["","**Progress:** "+" · ".join(f"{k}: {cnt.get(k,0)}" for k in ["Done","In review","In progress","Blocked","Not started","Deferred"])+f" (of {len(CORE)})"]
    active=[r for r in CORE if status_of(r[0])[0]!="Not started"]
    if active:
        out+=["","### Active and recently completed","","| ID | Work item | Status | Note |","|---|---|---|---|"]
        out+=[f"| {r[0]} | {r[2]} | {status_of(r[0])[0]} | {status_of(r[0])[1]} |" for r in active]
    out+=["","## Domain backlogs (detail, owned by their docs)",""]
    for k,v in DOMAIN.items():
        ids=[r[1] for r in v]
        out.append(f"- **{k}** — {len(v)} items ({ids[0] if ids else ''}…{ids[-1] if ids else ''}), source: `{v[0][0] if v else ''}`")
    out+=[f"- **Security tests** — {len(SECTESTS)} (06) · **Test matrix** — {len(TESTS)} (15)","",
          "## Risks","","| ID | Risk | Severity | Likelihood | Mitigation | Owner |","|---|---|---|---|---|---|"]
    out+= [f"| {' | '.join(r)} |" for r in RISKS]
    out+=["","## Decisions (ADRs)","","| ADR | Decision | Status |","|---|---|---|"]
    out+= [f"| [{d[0]}](adr/{d[3]}) | {d[1]} | {d[2]} |" for d in DECISIONS]
    out+=["","## Open questions","","See [22 §1](22-UNKNOWNS-ASSUMPTIONS-AND-OPEN-QUESTIONS.md). Decide-by milestones are tracked in the spreadsheet `Open Questions` sheet.","",
          "## RACI (delivery-level)","","| "+" | ".join(RACI[0])+" |","|"+"---|"*len(RACI[0])]
    out+= ["| "+" | ".join(r)+" |" for r in RACI[1:]]
    out+=["","## Cadence","","- Weekly: tracker review (status, blockers). - Per milestone: exit checklist review ([07 §3](07-ROADMAP-AND-MILESTONES.md)), register review ([22](22-UNKNOWNS-ASSUMPTIONS-AND-OPEN-QUESTIONS.md)), risk re-score, usability round.",""]
    (DOCS/"17-EXECUTION-TRACKER.md").write_text("\n".join(out))

HDR=Font(bold=True,color="FFFFFF"); FILL=PatternFill("solid",fgColor="1F2937")
def sheet(wb,name,header,rows,widths=None):
    ws=wb.create_sheet(name[:31]); ws.append(header)
    for r in rows: ws.append(list(r)+[""]*(len(header)-len(r)) if len(r)<len(header) else list(r)[:len(header)])
    for c in ws[1]: c.font=HDR; c.fill=FILL; c.alignment=Alignment(vertical="center")
    ws.freeze_panes="A2"
    for i,w in enumerate(widths or [14]*len(header)): ws.column_dimensions[chr(65+i) if i<26 else "A"+chr(65+i-26)].width=w
    if rows:
        ref=f"A1:{ws.cell(row=1,column=len(header)).column_letter}{len(rows)+1}"
        t=Table(displayName=re.sub(r'\W','',name)[:20]+"T",ref=ref); t.tableStyleInfo=TableStyleInfo(name="TableStyleLight9",showRowStripes=True); ws.add_table(t)
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment=Alignment(wrap_text=True,vertical="top")
    return ws

def write_xlsx():
    wb=Workbook(); ov=wb.active; ov.title="Overview"
    ov.append(["Procera — Delivery Tracker (generated 2026-10-11 by scripts/build_tracker.py)"]); ov["A1"].font=Font(bold=True,size=14)
    ov.append([]); ov.append(["Sheet","Contents"])
    for s,c in [("Milestones","Phased releases P0, R1–R5, GA, POST with dates and gates"),("Backlog","Core epics BPMS-xxx with status"),("Domain Backlog","TPL/UX/INT/CFG/AI items parsed from docs"),("Risks","Risk register"),("Decisions","ADR status"),("Security Tests","SEC-xx from 06"),("Test Matrix","T01–T40 from 15"),("Open Questions","OQ-xx"),("Unknowns","U-xx"),("Assumptions","A-xx"),("RACI","Delivery RACI"),("Burndown","Count of epics by milestone and status (formulas)")]:
        ov.append([s,c])
    ov.column_dimensions["A"].width=20; ov.column_dimensions["B"].width=70
    sheet(wb,"Milestones",["M","Name","Start","End","Gate","Scope","Status"],[m+("Not started",) for m in MILESTONES],[7,32,12,12,8,70,14])
    ws=sheet(wb,"Backlog",["ID","Area","Work item","Pri","Release","Owner role","Refs","Acceptance","Size","Depends","Status","Assignee","Start","Due","Notes"],
             [r+(status_of(r[0])[0],"","","",status_of(r[0])[1]) for r in CORE],[10,8,60,5,8,14,26,34,6,16,13,14,11,11,30])
    from openpyxl.worksheet.datavalidation import DataValidation
    dv=DataValidation(type="list",formula1='"Not started,In progress,Blocked,In review,Done,Deferred"',allow_blank=True); ws.add_data_validation(dv); dv.add(f"K2:K{len(CORE)+1}")
    red=PatternFill("solid",fgColor="FECACA"); green=PatternFill("solid",fgColor="BBF7D0"); amber=PatternFill("solid",fgColor="FDE68A")
    ws.conditional_formatting.add(f"K2:K{len(CORE)+1}",CellIsRule(operator="equal",formula=['"Blocked"'],fill=red))
    ws.conditional_formatting.add(f"K2:K{len(CORE)+1}",CellIsRule(operator="equal",formula=['"Done"'],fill=green))
    ws.conditional_formatting.add(f"K2:K{len(CORE)+1}",CellIsRule(operator="equal",formula=['"In progress"'],fill=amber))
    drows=[]
    for k,v in DOMAIN.items():
        for r in v: drows.append([k, r[1], " | ".join(r[2:]), r[0], "Not started"])
    sheet(wb,"Domain Backlog",["Domain","ID","Details (as in source doc)","Source doc","Status"],drows,[18,10,100,30,13])
    sheet(wb,"Risks",["ID","Risk","Severity","Likelihood","Mitigation","Owner","Trend"],[r+("—",) for r in RISKS],[6,36,10,10,60,14,8])
    sheet(wb,"Decisions",["ADR","Decision","Status","File"],DECISIONS,[10,60,36,50])
    sheet(wb,"Security Tests",["Source","ID","Test","Pass condition","Phase/M","Extra"],SECTESTS,[26,8,50,50,14,20])
    sheet(wb,"Test Matrix",["Source","ID","Test","Category","Pass condition","Milestone","Status"],TESTS,[26,6,40,14,50,10,12])
    sheet(wb,"Open Questions",["Source","ID","Question","Recommendation","Owner","Decide by","Decision"],[r+[""] for r in OQS],[20,7,50,60,18,10,30])
    sheet(wb,"Unknowns",["Source","ID","Unknown","Why it matters","Resolution","Decide by"],UNK,[20,6,50,40,50,10])
    sheet(wb,"Assumptions",["Source","ID","Assumption","Check","If false"],ASM,[20,6,60,40,40])
    sheet(wb,"RACI",list(RACI[0]),RACI[1:],[32,14,12,18,16,16])
    bd=wb.create_sheet("Burndown"); stats=["Not started","In progress","Blocked","In review","Done","Deferred"]
    bd.append(["Release"]+stats+["Total"])
    for m in [x[0] for x in MILESTONES]:
        r=bd.max_row+1
        bd.append([m]+[f'=COUNTIFS(Backlog!$E:$E,$A{r},Backlog!$K:$K,"{s}")' for s in stats]+[f"=COUNTIF(Backlog!$E:$E,$A{r})"])
    for c in bd[1]: c.font=HDR; c.fill=FILL
    out=ROOT/"trackers"; out.mkdir(exist_ok=True); wb.save(out/"Delivery_Tracker.xlsx")
    with open(out/"backlog.csv","w",newline="") as f:
        w=csv.writer(f); w.writerow(["ID","Area","Work item","Pri","Release","Owner","Refs","Acceptance","Size","Depends","Status"]); [w.writerow(list(r)+[status_of(r[0])[0]]) for r in CORE]
    with open(out/"domain_backlog.csv","w",newline="") as f:
        w=csv.writer(f); w.writerow(["Domain","ID","Details","Source"]); [w.writerow(r[:4]) for r in drows]

def bar(done, total, width=20):
    if total == 0: return "░" * width
    n = round(width * done / total)
    return "█" * n + "░" * (width - n)

def e2e_count():
    """E2E scenarios implemented as Playwright tests (ids in test titles)."""
    import re, pathlib
    ids = set()
    for f in pathlib.Path("e2e/packs").rglob("*.spec.ts"):
        ids |= set(re.findall(r"test\(\s*'(E2E-[A-Z0-9-]+)", f.read_text()))
    return len(ids)

def rust_test_count():
    import re, pathlib
    n = 0
    for f in list(pathlib.Path("crates").rglob("*.rs")) + list(pathlib.Path("apps").rglob("*.rs")):
        n += len(re.findall(r"#\[(?:tokio::)?test\]", f.read_text()))
    return n

def recent_commits(n):
    import subprocess
    try:
        out = subprocess.run(["git", "log", f"-{n}", "--format=%h\t%s"], capture_output=True, text=True, check=True).stdout
    except Exception:
        return []
    return [tuple(l.split("\t", 1)) for l in out.splitlines() if "\t" in l]

def write_dashboard():
    from collections import Counter, defaultdict
    by_rel = defaultdict(list)
    for r in CORE: by_rel[r[4]].append(r)
    cnt = Counter(status_of(r[0])[0] for r in CORE)
    done = cnt.get("Done", 0); total = len(CORE)
    current = next((m for m in MILESTONES if any(status_of(r[0])[0] != "Done" for r in by_rel[m[0]])), MILESTONES[-1])
    o = ["# Procera — tracker dashboard", "",
         "> **Generated** by `scripts/build_tracker.py`. Don't edit by hand. To change status, edit the `STATUS` map (or `CORE`) in the script and re-run it ([rules](../docs/24-ENGINEERING-RULES.md#2-tracker-rules)).",
         "> Detail: [docs/17](../docs/17-EXECUTION-TRACKER.md) · [Delivery_Tracker.xlsx](Delivery_Tracker.xlsx) · [backlog.csv](backlog.csv) · [domain_backlog.csv](domain_backlog.csv) · Roadmap: [docs/07](../docs/07-ROADMAP-AND-MILESTONES.md)", "",
         "## At a glance", "",
         f"```text",
         f"Overall   {bar(done,total,30)}  {done}/{total} epics done  ({round(100*done/total)}%)",
         f"Partial   {bar(done + 0.5*cnt.get('In progress',0), total, 30)}  {round(100*(done + 0.5*cnt.get('In progress',0))/total)}% counting in-progress epics as half done",
         f"Tests     {e2e_count()} automated E2E scenarios · {rust_test_count()} Rust tests (incl. property tests)",
         f"Status    done {done} · in review {cnt.get('In review',0)} · in progress {cnt.get('In progress',0)} · blocked {cnt.get('Blocked',0)} · not started {cnt.get('Not started',0)} · deferred {cnt.get('Deferred',0)}",
         f"Now       {current[0]} — {current[1]}  ({current[2]} → {current[3]}, gate {current[4]})",
         "```", "", "## Releases", "", "| Release | Name | Window | Gate | Progress | Done |", "|---|---|---|---|---|---|"]
    for m in MILESTONES:
        items = by_rel[m[0]]; d = sum(1 for r in items if status_of(r[0])[0] == "Done")
        marker = " ◀ now" if m[0] == current[0] else ""
        o.append(f"| **{m[0]}**{marker} | {m[1]} | {m[2]} → {m[3]} | {m[4]} | `{bar(d, len(items), 12)}` | {d}/{len(items)} |")
    active = [r for r in CORE if status_of(r[0])[0] in ("In progress", "In review", "Blocked")]
    o += ["", "## Active right now", ""]
    if active:
        o += ["| ID | Item | Status | Note |", "|---|---|---|---|"]
        o += [f"| {r[0]} | {r[2]} | **{status_of(r[0])[0]}** | {status_of(r[0])[1]} |" for r in active]
    else:
        o.append("_Nothing in progress._")
    shipped = recent_commits(12)
    if shipped:
        o += ["", "## Recently shipped (last 12 commits on main)", ""] + [f"- `{h}` {t}" for h, t in shipped]
    recent = [r for r in CORE if status_of(r[0])[0] == "Done"]
    o += ["", "## Done", ""]
    o += [f"- [x] **{r[0]}** {r[2]} — {status_of(r[0])[1]}" for r in recent] or ["_Nothing done yet._"]
    o += ["", "## Release checklists", "",
          "Legend: `[x]` done · `[ ]` open. The status is shown in brackets when the item isn't simply not started.", ""]
    for m in MILESTONES:
        items = by_rel[m[0]]
        if not items: continue
        d = sum(1 for r in items if status_of(r[0])[0] == "Done")
        o += [f"<details{' open' if m[0]==current[0] else ''}>",
              f"<summary><b>{m[0]} — {m[1]}</b> · {m[2]} → {m[3]} · gate {m[4]} · {d}/{len(items)} done</summary>", "",
              f"_{m[5]}_", ""]
        for r in items:
            st = status_of(r[0])[0]
            box = "x" if st == "Done" else " "
            tag = "" if st in ("Done", "Not started") else f" **[{st}]**"
            o.append(f"- [{box}] **{r[0]}** ({r[3]}, {r[8]}) {r[2]}{tag}")
        o += ["", "</details>", ""]
    o += ["## Domain backlogs", "", "| Backlog | Items | Source |", "|---|---|---|"]
    for k, v in DOMAIN.items():
        o.append(f"| {k} | {len(v)} | [{v[0][0] if v else ''}](../docs/{v[0][0] if v else ''}) |")
    o += [f"| Security tests | {len(SECTESTS)} | [06](../docs/06-SECURITY-AND-OPERATIONS.md) |",
          f"| Test matrix | {len(TESTS)} | [15](../docs/15-TEST-STRATEGY-AND-CONFORMANCE.md) |",
          f"| Open questions / unknowns / assumptions | {len(OQS)} / {len(UNK)} / {len(ASM)} | [22](../docs/22-UNKNOWNS-ASSUMPTIONS-AND-OPEN-QUESTIONS.md) |", ""]
    (ROOT / "trackers" / "README.md").write_text("\n".join(o))

write_md(); write_xlsx(); write_dashboard()
print("core",len(CORE),"domain",{k:len(v) for k,v in DOMAIN.items()},"sec",len(SECTESTS),"tests",len(TESTS),"oq",len(OQS),"unk",len(UNK),"asm",len(ASM),"adr",len(DECISIONS))
